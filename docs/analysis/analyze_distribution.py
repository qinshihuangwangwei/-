#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
原始综测成绩表列取值分布分析

用途：从真实的 `25级综测成绩*.xlsx` 中识别"基线制"计分项
（即某个子项在全表所有学生中取值恒定 —— 说明它不是"不申请就 0 分"的
申请制，而是"默认给分、只录扣分"的基线制）。

这一步发现的问题直接决定了 ScoringItem 需要 `计分模式` 字段。
详见 docs/superpowers/specs/2026-09-26-zongce-scoring-system-design.md §5.6

用法：
    python analyze_distribution.py <xlsx路径> [输出txt路径]
"""

import sys
import io
from collections import Counter

import openpyxl

# 列号(1-based) -> 计分项名称。依据原表第 2/3/4 行的多级合并表头人工核对得出。
COLUMN_NAMES = {
    4:  "A1-1 志愿服务立项及开展(4)",
    5:  "A1-2 日常志愿服务及社会服务奉献(2)",
    6:  "A2-1 学生干部任职加分(6)",
    7:  "A2-2 参与学生工作-主管部门评定(6)",
    8:  "A2-3 参与学生工作-班级评定(4)",
    9:  "A2-4 劳动(2)",
    10: "A3-1 党建活动(3)",
    11: "A3-2 团建活动-班团活动(5) [K:L 合并]",
    12: "  (K:L 合并的从属格，恒空)",
    13: "A3-3 团建活动-青年大学习(3)",
    14: "A3-4 团建活动-突出贡献(2)",
    15: "A4-1 寝室检查(12)",
    16: "A4-2 校院级公寓文化节活动(2)",
    17: "A4-3 星级寝室评比(2)",
    18: "A5-1 立项加分(2)",
    19: "A5-2 社会实践报告-团体(3)",
    20: "A5-3 社会实践报告-个人(3)",
    21: "A6 遵纪守法(6) [U:W 合并]",
    22: "  (U:W 合并的从属格，恒空)",
    23: "  (U:W 合并的从属格，恒空)",
    24: "A7-1 社团(3)",
    25: "A7-2 文艺类(5)",
    26: "A7-3 体育类(6)",
    27: "A7-4 创新创业类(3)",
    28: "A7-5 其他素质能力(8)",
    29: "A8-1 易班-个人活跃度(2)",
    30: "A8-2 易班APP线上活动(4)",
    31: "A8-3 易班线下活动(2)",
    32: "A 项合计(100)",
}


def analyze(path, out_path):
    wb = openpyxl.load_workbook(path, data_only=True)
    out = io.open(out_path, "w", encoding="utf-8")
    findings = []

    for ws in wb.worksheets:
        out.write("\n===== %s =====\n" % ws.title)

        # 数据行从第 5 行开始（第 1 行标题，第 2-4 行多级表头）
        rows = [
            r for r in range(5, ws.max_row + 1)
            if ws.cell(r, 1).value not in (None, "")
        ]
        if not rows:
            continue
        out.write("学生数: %d\n" % len(rows))

        for c in range(4, 32):
            vals = [ws.cell(r, c).value for r in rows]
            nonnull = [v for v in vals if v not in (None, "")]
            name = COLUMN_NAMES.get(c, "col%d" % c)

            if not nonnull:
                out.write("  col%2d %-42s 全空\n" % (c, name))
                continue

            cnt = Counter(str(v) for v in nonnull)
            top, top_n = cnt.most_common(1)[0]
            distinct = len(cnt)

            tag = ""
            if distinct == 1:
                tag = "  <<< 全表同一个值"
                findings.append((ws.title, name, top, len(nonnull)))
            elif top_n / len(nonnull) >= 0.9:
                tag = "  <<< %.0f%% 都是 %s" % (100 * top_n / len(nonnull), top)

            out.write(
                "  col%2d %-42s 非空%3d 取值%2d 最常见=%s(%d)%s\n"
                % (c, name, len(nonnull), distinct, top, top_n, tag)
            )

    out.write("\n\n===== 恒定值汇总（候选项：基线制计分项） =====\n")
    for sheet, name, val, n in findings:
        out.write("  %-24s %-42s 恒为 %s (%d人)\n" % (sheet, name, val, n))

    out.close()
    return findings


if __name__ == "__main__":
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(1)
    src = sys.argv[1]
    dst = sys.argv[2] if len(sys.argv) > 2 else "distribution-analysis.txt"
    result = analyze(src, dst)
    print("分析完成，恒定值条目 %d 个，输出：%s" % (len(result), dst))
